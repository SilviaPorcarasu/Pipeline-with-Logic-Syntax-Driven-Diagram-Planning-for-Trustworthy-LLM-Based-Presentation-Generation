"""PowerPoint rendering helpers built on python-pptx."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt, Emu

from layout_engine import get_layout
from models import SlideContent


PLAIN_THEME: dict[str, object] = {
    "bg": (255, 255, 255),
    "title": (25, 45, 84),
    "subtitle": (80, 100, 130),
    "body": (40, 50, 70),
    "accent": (0, 90, 180),
    "title_font": "Times New Roman",
    "subtitle_font": "Times New Roman",
    "body_font": "Times New Roman",
}


def _set_background(slide, *, rgb: tuple[int, int, int]) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(*rgb)


def _style_runs(
    paragraph,
    *,
    size: int,
    color: tuple[int, int, int],
    bold: bool = False,
    italic: bool = False,
    font_name: str = "Times New Roman",
) -> None:
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(*color)
        run.font.bold = bool(bold)
        run.font.italic = bool(italic)
        run.font.name = font_name


def _clear_shape_chrome(shape) -> None:
    """Remove shape fill/outline so containers are invisible."""
    try:
        shape.fill.background()
    except Exception:
        pass
    try:
        shape.line.fill.background()
    except Exception:
        pass
    try:
        shape.line.width = Pt(0)
    except Exception:
        pass


def _add_title_underline(slide, *, x: float, y: float, w: float, color: tuple[int, int, int]) -> None:
    """Add a thin accent line under the title."""
    line = slide.shapes.add_shape(
        1,  # MSO_SHAPE.RECTANGLE
        Inches(x + 0.5), Inches(y), Inches(w - 1.0), Emu(28000),
    )
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(*color)
    try:
        line.line.fill.background()
    except Exception:
        pass


def _pick_title_font_size(text: str, *, box_w: float, box_h: float) -> int:
    normalized = " ".join((text or "").split())
    length = len(normalized)
    longest_word = max((len(part) for part in normalized.split()), default=0)
    if length <= 22:
        size = 36
    elif length <= 32:
        size = 32
    elif length <= 44:
        size = 28
    elif length <= 58:
        size = 26
    else:
        size = 24
    if longest_word >= 14:
        size -= 2
    if box_w < 10.0:
        size -= 1
    return max(20, size)


def _add_title(
    slide,
    text: str,
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    theme: dict[str, object],
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _clear_shape_chrome(box)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    p.space_after = Pt(4)
    _style_runs(
        p,
        size=_pick_title_font_size(text, box_w=w, box_h=h),
        color=theme["title"],
        bold=True,
        font_name=str(theme["title_font"]),
    )


def _add_subtitle(
    slide,
    text: str,
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    theme: dict[str, object],
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _clear_shape_chrome(box)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.LEFT
    _style_runs(
        p,
        size=24,
        color=theme["subtitle"],
        bold=False,
        italic=True,
        font_name=str(theme["subtitle_font"]),
    )


def _add_bullets(
    slide,
    bullets: list[str],
    *,
    x: float,
    y: float,
    w: float,
    h: float,
    text_only: bool = False,
    theme: dict[str, object],
) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    _clear_shape_chrome(box)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    # Font size based on number of bullets and mode
    if text_only:
        size = 24
    elif len(bullets) <= 2:
        size = 24
    elif len(bullets) <= 3:
        size = 22
    else:
        size = 20
    for idx, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        if text_only:
            p.text = bullet
            p.alignment = PP_ALIGN.CENTER
        else:
            p.text = f"\u2022  {bullet}"
            p.alignment = PP_ALIGN.LEFT
        # Spacing between bullets
        p.space_before = Pt(6) if idx > 0 else Pt(0)
        p.space_after = Pt(8)
        _style_runs(
            p,
            size=size,
            color=theme["body"],
            bold=False,
            font_name=str(theme["body_font"]),
        )


def _add_image(slide, image_path: Path, *, x: float, y: float, w: float, h: float) -> None:
    pic = slide.shapes.add_picture(str(image_path), Inches(x), Inches(y), width=Inches(w), height=Inches(h))
    _clear_shape_chrome(pic)


def render_ppt(
    *,
    out_path: Path,
    slides: list[SlideContent],
    layout_by_slide: dict[int, str],
    image_by_slide: dict[int, Path],
    deck_style: str = "academic",
) -> Path:
    """Render the final PowerPoint deck."""
    out_path.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    theme = PLAIN_THEME

    ordered = sorted(slides, key=lambda s: s.slide_number)
    for slide_content in ordered:
        layout_name = layout_by_slide.get(slide_content.slide_number, slide_content.layout_hint or "text_only")
        layout = get_layout(layout_name)

        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, rgb=theme["bg"])

        # Title
        _add_title(
            slide,
            slide_content.title,
            x=layout.title_x,
            y=layout.title_y,
            w=layout.title_w,
            h=layout.title_h,
            theme=theme,
        )

        # Accent line under title
        if layout_name not in {"title", "section_divider"}:
            _add_title_underline(
                slide,
                x=layout.title_x,
                y=layout.title_y + layout.title_h + 0.05,
                w=layout.title_w,
                color=theme["accent"],
            )

        # Subtitle
        subtitle_text = " ".join((slide_content.subtitle or "").split())
        body_y = layout.body_y
        body_h = layout.body_h
        if subtitle_text:
            subtitle_y = layout.title_y + layout.title_h + 0.15
            subtitle_h = 0.45
            _add_subtitle(
                slide,
                subtitle_text,
                x=layout.title_x + 0.05,
                y=subtitle_y,
                w=layout.title_w,
                h=subtitle_h,
                theme=theme,
            )
            body_y = subtitle_y + subtitle_h + 0.25
            body_h = max(0.8, (layout.body_y + layout.body_h) - body_y)

        # Bullets
        is_title_like = layout_name in {"title", "section_divider", "conclusion"}
        _add_bullets(
            slide,
            slide_content.bullets,
            x=layout.body_x,
            y=body_y,
            w=layout.body_w,
            h=body_h,
            text_only=is_title_like,
            theme=theme,
        )

        # Image / diagram
        image_path = image_by_slide.get(slide_content.slide_number)
        if (
            image_path is not None
            and image_path.exists()
            and layout.image_x is not None
            and layout.image_y is not None
            and layout.image_w is not None
            and layout.image_h is not None
        ):
            _add_image(
                slide,
                image_path,
                x=layout.image_x,
                y=layout.image_y,
                w=layout.image_w,
                h=layout.image_h,
            )

        # Voiceover in notes
        notes = slide.notes_slide.notes_text_frame
        notes.clear()
        notes.text = slide_content.voiceover.strip()

    prs.save(str(out_path))
    return out_path
